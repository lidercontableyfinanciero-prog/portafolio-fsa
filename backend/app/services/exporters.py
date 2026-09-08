"""Exportadores de la vista actual (Dashboard y tablas filtradas) a Excel y PDF."""

from __future__ import annotations

import io
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Image,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

LOGO_PATH = Path(__file__).resolve().parents[1] / "assets" / "logo_fsa.png"


def _logo_flowable(width_mm: float = 34):
    """Logo institucional para la esquina superior izquierda del PDF."""
    if not LOGO_PATH.exists():
        return None
    try:
        img = Image(str(LOGO_PATH))
        ratio = img.imageHeight / img.imageWidth
        img.drawWidth = width_mm * mm
        img.drawHeight = width_mm * mm * ratio
        img.hAlign = "LEFT"
        return img
    except Exception:  # noqa: BLE001
        return None

NAVY = colors.HexColor("#0E2841")
GREY = colors.HexColor("#6B7280")
BORDER = colors.HexColor("#E2E5EA")
GREEN = colors.HexColor("#1E7B34")
RED = colors.HexColor("#C0392B")

_HEADER_FILL = PatternFill("solid", fgColor="0E2841")
_HEADER_FONT = Font(color="FFFFFF", bold=True)

POSITION_COLUMNS: list[tuple[str, str]] = [
    ("description", "Descripción"),
    ("identifier", "CUSIP/CINS"),
    ("classification", "Clasificación"),
    ("type", "Tipo"),
    ("sector", "Sector"),
    ("cost_basis", "Costo (USD)"),
    ("market_value", "Valor de Mercado (USD)"),
    ("unrealized_gain_loss", "G/(P) No Realizada (USD)"),
    ("return_on_cost", "Rentab. s/ Costo"),
    ("annual_income", "Ingreso Anual Est. (USD)"),
    ("accrued_interest", "Interés Acumulado (USD)"),
    ("dividends_paid", "Intereses/Dividendos Pagados (USD)"),
    ("tax", "Impuesto (USD)"),
    ("tax_rate", "Tasa Impositiva"),
    ("current_yield", "Yield Actual"),
    ("equity_return_on_cost", "Rentab. Costo (Renta Var.)"),
    ("equity_market_value_return", "Rentab. Valor Mercado (Renta Var.)"),
    ("moodys_rating", "Moody's"),
    ("moodys_grade", "Grado Moody's"),
    ("sp_rating", "S&P"),
    ("sp_grade", "Grado S&P"),
    ("stop_loss", "Indicador Stop-Loss"),
    ("time_alert", "Alerta Tiempo"),
    ("issuer_alert", "Alerta Emisor"),
]
_PCT_COLS = {
    "return_on_cost", "tax_rate", "current_yield",
    "equity_return_on_cost", "equity_market_value_return",
}
_MONEY_COLS = {
    "cost_basis", "market_value", "unrealized_gain_loss", "annual_income",
    "accrued_interest", "dividends_paid", "tax",
}

_MONEY = "#,##0.00"
_PCT = "0.00%"


def _meta_line(meta: dict) -> str:
    p = meta.get("period", {})
    parts = [f"Período: {p.get('month', '?')} {p.get('year', '?')}"]
    f = meta.get("filters") or {}
    applied = [f"{k}={v}" for k, v in f.items() if v]
    if applied:
        parts.append("Filtros: " + ", ".join(applied))
    parts.append(f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    return " · ".join(parts)


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
def _style_header(ws, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")


def _autosize(ws, max_width: int = 48) -> None:
    for col in ws.columns:
        length = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(length + 2, max_width)


def positions_to_xlsx(rows: list[dict], meta: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Posiciones"
    ws.append([label for _, label in POSITION_COLUMNS])
    _style_header(ws, len(POSITION_COLUMNS))
    for r in rows:
        ws.append([r.get(key) for key, _ in POSITION_COLUMNS])
    for idx, (key, _) in enumerate(POSITION_COLUMNS, start=1):
        fmt = _MONEY if key in _MONEY_COLS else _PCT if key in _PCT_COLS else None
        if fmt:
            for cell in ws[get_column_letter(idx)][1:]:
                cell.number_format = fmt
    ws.freeze_panes = "A2"
    _autosize(ws)

    info = wb.create_sheet("Info")
    info["A1"] = "Portafolio FSA — Posiciones"
    info["A2"] = _meta_line(meta)
    info["A3"] = f"Total de posiciones exportadas: {len(rows)}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def dashboard_to_xlsx(dash: dict, evolution: list[dict], twr: dict, meta: dict) -> bytes:
    wb = Workbook()

    k = dash["kpis"]
    ws = wb.active
    ws.title = "Indicadores"
    ws.append(["Portafolio FSA — Indicadores del período"])
    ws.append([_meta_line(meta)])
    ws.append([])
    kpi_rows = [
        ("Costo Total (USD)", k["costo_total"], _MONEY),
        ("Valor de Mercado (USD)", k["valor_mercado"], _MONEY),
        ("G/(P) No Realizada (USD)", k["gp_no_realizada"], _MONEY),
        ("Rentabilidad s/ Costo", k["rentab_sobre_costo"], _PCT),
        ("Ingreso Anual Estimado (USD)", k["ingreso_anual_est"], _MONEY),
        ("Interés Acumulado (USD)", k["interes_acumulado"], _MONEY),
        ("Yield Promedio Ponderado", k["yield_prom_ponderado"], _PCT),
        ("N° de Posiciones", k["n_posiciones"], "0"),
    ]
    for name, val, fmt in kpi_rows:
        ws.append([name, val])
        ws.cell(row=ws.max_row, column=2).number_format = fmt
    ws.append([])
    ws.append(["Rentabilidad del último mes (Dietz)", _last_return(twr)])
    ws.cell(row=ws.max_row, column=2).number_format = _PCT
    ws.append(["TWR acumulado del año", twr.get("cumulative_twr", 0)])
    ws.cell(row=ws.max_row, column=2).number_format = _PCT
    ws.append(["Benchmark acumulado", twr.get("cumulative_benchmark", 0)])
    ws.cell(row=ws.max_row, column=2).number_format = _PCT
    ws.append(["Alfa acumulado", twr.get("cumulative_twr", 0) - twr.get("cumulative_benchmark", 0)])
    ws.cell(row=ws.max_row, column=2).number_format = _PCT
    _autosize(ws)

    _breakdown_sheet(wb, "Por Clasificación", dash["por_clasificacion"])
    _breakdown_sheet(wb, "Por Tipo", dash["por_tipo"])
    _breakdown_sheet(wb, "Por Sector", dash["por_sector"])
    _risk_sheet(wb, dash)
    _evolution_sheet(wb, evolution)
    _twr_sheet(wb, twr)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _last_return(twr: dict) -> float:
    rows = [r for r in twr.get("rows", []) if r.get("portfolio_return") is not None]
    return rows[-1]["portfolio_return"] if rows else 0.0


def _breakdown_sheet(wb: Workbook, title: str, rows: list[dict]) -> None:
    ws = wb.create_sheet(title[:31])
    ws.append(["Categoría", "Costo", "Valor de Mercado", "G/(P) No Realizada",
               "% Participación", "Ingreso Anual Est.", "Posiciones"])
    _style_header(ws, 7)
    for r in rows:
        ws.append([r["label"], r["costo"], r["valor_mercado"], r["gp_no_realizada"],
                   r["pct_participacion"], r["ingreso_anual_est"], r["posiciones"]])
    for col in "BCDF":
        for cell in ws[col][1:]:
            cell.number_format = _MONEY
    for cell in ws["E"][1:]:
        cell.number_format = _PCT
    _autosize(ws)


def _risk_sheet(wb: Workbook, dash: dict) -> None:
    ws = wb.create_sheet("Riesgo y Alertas")
    for panel, title in (
        ("calidad_moodys", "Calidad crediticia (Moody's)"),
        ("calidad_sp", "Calidad crediticia (S&P)"),
        ("stop_loss", "Indicador Stop-Loss"),
        ("alerta_tiempo", "Alerta Tiempo (plazo de tenencia)"),
        ("alerta_emisor", "Alerta Emisor (concentración)"),
        ("limite_cash", "Límite de Caja"),
    ):
        ws.append([title])
        ws.append(["Categoría", "Posiciones", "Valor de Mercado", "% del Total"])
        for r in dash.get(panel, []):
            ws.append([r["label"], r["posiciones"], r["valor_mercado"], r["pct_participacion"]])
            ws.cell(row=ws.max_row, column=3).number_format = _MONEY
            ws.cell(row=ws.max_row, column=4).number_format = _PCT
        ws.append([])
    ra = dash.get("risk_alerts", {})
    ws.append(["Vencimientos < 1 año (posiciones)", ra.get("vencimientos_1a_posiciones", 0)])
    ws.append(["Vencimientos < 1 año (valor USD)", ra.get("vencimientos_1a_valor", 0)])
    ws.cell(row=ws.max_row, column=2).number_format = _MONEY
    ws.append(["Plazo prom. al vencimiento — bonos (años)", ra.get("plazo_prom_vencimiento_bonos", 0)])
    ws.append(["Emisores sobre el límite de 500k USD", ra.get("emisores_sobre_limite", 0)])
    _autosize(ws)


def _evolution_sheet(wb: Workbook, evolution: list[dict]) -> None:
    ws = wb.create_sheet("Evolución Mensual")
    ws.append(["Mes", "Costo", "Valor de Mercado", "G/(P) No Realizada",
               "Ingreso Anual Est.", "Rentab. s/ Costo"])
    _style_header(ws, 6)
    for p in evolution:
        ws.append([p["label"], p["costo"], p["valor_mercado"], p["gp_no_realizada"],
                   p["ingreso_anual_est"], p["rentab_sobre_costo"]])
    for col in "BCDE":
        for cell in ws[col][1:]:
            cell.number_format = _MONEY
    for cell in ws["F"][1:]:
        cell.number_format = _PCT
    _autosize(ws)


def _twr_sheet(wb: Workbook, twr: dict) -> None:
    ws = wb.create_sheet("Rentabilidad TWR")
    ws.append(["Mes", "R mes (Dietz)", "Benchmark mes", "Factor (1+R)", "TWR acumulado"])
    _style_header(ws, 5)
    for r in twr.get("rows", []):
        ws.append([f"{r['month_name']} {r['year']}", r.get("portfolio_return"),
                   r.get("benchmark_return"), r.get("factor"), r.get("cumulative_twr")])
    for col in "BCE":
        for cell in ws[col][1:]:
            cell.number_format = _PCT
    _autosize(ws)


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
def _pdf_styles():
    ss = getSampleStyleSheet()
    ss.add(ParagraphStyle("FSATitle", parent=ss["Title"], textColor=NAVY, fontSize=16))
    ss.add(ParagraphStyle("FSAMeta", parent=ss["Normal"], textColor=GREY, fontSize=8))
    ss.add(ParagraphStyle("FSAH2", parent=ss["Heading2"], textColor=NAVY, fontSize=11,
                          spaceBefore=10, spaceAfter=4))
    return ss


def _base_table_style() -> TableStyle:
    return TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), NAVY),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), 0.4, BORDER),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F7F9")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ])


def _fmt_money(v) -> str:
    return "—" if v is None else f"{v:,.0f}"


def _fmt_pct(v) -> str:
    return "—" if v is None else f"{v * 100:,.2f}%"


def positions_to_pdf(rows: list[dict], meta: dict) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
        title="Portafolio FSA — Posiciones",
    )
    ss = _pdf_styles()
    story = []
    logo = _logo_flowable()
    if logo is not None:
        story += [logo, Spacer(1, 4)]
    story += [
        Paragraph("Portafolio FSA — Posiciones", ss["FSATitle"]),
        Paragraph(_meta_line(meta), ss["FSAMeta"]),
        Spacer(1, 6),
    ]
    head = ["Descripción", "CUSIP", "Clas.", "Tipo", "Sector", "Costo", "V. Mercado",
            "G/(P)", "Rent.", "Stop-Loss", "Moody's"]
    data = [head]
    for r in rows:
        data.append([
            (r.get("description") or "")[:28],
            "—" if r.get("identifier") == r.get("description") else (r.get("identifier") or "")[:12],
            (r.get("classification") or "")[:10],
            r.get("type") or "",
            (r.get("sector") or "")[:16],
            _fmt_money(r.get("cost_basis")),
            _fmt_money(r.get("market_value")),
            _fmt_money(r.get("unrealized_gain_loss")),
            _fmt_pct(r.get("return_on_cost")),
            (r.get("stop_loss") or "")[:22],
            (r.get("moodys_grade") or "")[:10],
        ])
    t = Table(data, repeatRows=1,
              colWidths=[52 * mm, 22 * mm, 18 * mm, 14 * mm, 30 * mm, 20 * mm, 22 * mm,
                         20 * mm, 16 * mm, 34 * mm, 22 * mm])
    tstyle = _base_table_style()
    for i, r in enumerate(rows, start=1):
        gl = r.get("unrealized_gain_loss") or 0
        tstyle.add("TEXTCOLOR", (7, i), (7, i), GREEN if gl >= 0 else RED)
    t.setStyle(tstyle)
    story.append(t)
    story.append(Spacer(1, 6))
    story.append(Paragraph(f"{len(rows)} posiciones · cifras en USD", ss["FSAMeta"]))
    doc.build(story)
    return buf.getvalue()


def dashboard_to_pdf(dash: dict, evolution: list[dict], twr: dict, meta: dict) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=15 * mm, rightMargin=15 * mm, topMargin=14 * mm, bottomMargin=14 * mm,
        title="Portafolio FSA — Dashboard",
    )
    ss = _pdf_styles()
    k = dash["kpis"]
    alfa = twr.get("cumulative_twr", 0) - twr.get("cumulative_benchmark", 0)
    story = []
    logo = _logo_flowable()
    if logo is not None:
        story += [logo, Spacer(1, 4)]
    story += [
        Paragraph("Portafolio FSA — Resumen del Dashboard", ss["FSATitle"]),
        Paragraph(_meta_line(meta), ss["FSAMeta"]),
        Spacer(1, 8),
        Paragraph("Indicadores del período", ss["FSAH2"]),
        _kv_table([
            ("Costo Total", _fmt_money(k["costo_total"]) + " USD"),
            ("Valor de Mercado", _fmt_money(k["valor_mercado"]) + " USD"),
            ("G/(P) No Realizada", _fmt_money(k["gp_no_realizada"]) + " USD"),
            ("Rentabilidad s/ Costo", _fmt_pct(k["rentab_sobre_costo"])),
            ("Ingreso Anual Estimado", _fmt_money(k["ingreso_anual_est"]) + " USD"),
            ("Interés Acumulado", _fmt_money(k["interes_acumulado"]) + " USD"),
            ("Yield Promedio Ponderado", _fmt_pct(k["yield_prom_ponderado"])),
            ("N° de Posiciones", str(k["n_posiciones"])),
        ]),
        Spacer(1, 6),
        Paragraph("Rentabilidad", ss["FSAH2"]),
        _kv_table([
            ("Rentabilidad del último mes (Dietz)", _fmt_pct(_last_return(twr))),
            ("TWR acumulado del año", _fmt_pct(twr.get("cumulative_twr", 0))),
            ("Benchmark acumulado", _fmt_pct(twr.get("cumulative_benchmark", 0))),
            ("Alfa acumulado", _fmt_pct(alfa)),
        ]),
        Spacer(1, 6),
        Paragraph("Distribución por clasificación", ss["FSAH2"]),
        _breakdown_pdf_table(dash["por_clasificacion"]),
        Spacer(1, 6),
        Paragraph("Distribución por tipo de instrumento", ss["FSAH2"]),
        _breakdown_pdf_table(dash["por_tipo"]),
        Spacer(1, 6),
        Paragraph("Riesgo y alertas", ss["FSAH2"]),
        _risk_pdf_table(dash),
        Spacer(1, 6),
        Paragraph("Rentabilidad acumulada (TWR) vs Benchmark", ss["FSAH2"]),
        _twr_pdf_table(twr),
    ]
    doc.build(story)
    return buf.getvalue()


def _kv_table(pairs: list[tuple[str, str]]) -> Table:
    t = Table([[k, v] for k, v in pairs], colWidths=[85 * mm, 85 * mm])
    t.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("TEXTCOLOR", (0, 0), (0, -1), GREY),
        ("TEXTCOLOR", (1, 0), (1, -1), NAVY),
        ("FONTNAME", (1, 0), (1, -1), "Helvetica-Bold"),
        ("LINEBELOW", (0, 0), (-1, -1), 0.3, BORDER),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    return t


def _breakdown_pdf_table(rows: list[dict]) -> Table:
    data = [["Categoría", "Costo", "V. Mercado", "G/(P)", "%", "Pos."]]
    for r in rows:
        data.append([
            r["label"], _fmt_money(r["costo"]), _fmt_money(r["valor_mercado"]),
            _fmt_money(r["gp_no_realizada"]), _fmt_pct(r["pct_participacion"]),
            str(r["posiciones"]),
        ])
    t = Table(data, repeatRows=1,
              colWidths=[55 * mm, 25 * mm, 28 * mm, 25 * mm, 18 * mm, 15 * mm])
    t.setStyle(_base_table_style())
    return t


def _risk_pdf_table(dash: dict) -> Table:
    data = [["Panel", "Categoría", "Pos.", "Valor de Mercado", "% Total"]]
    for panel, title in (
        ("calidad_moodys", "Moody's"),
        ("calidad_sp", "S&P"),
        ("stop_loss", "Stop-Loss"),
        ("alerta_tiempo", "Alerta Tiempo"),
        ("alerta_emisor", "Alerta Emisor"),
    ):
        for r in dash.get(panel, []):
            data.append([title, r["label"], str(r["posiciones"]),
                         _fmt_money(r["valor_mercado"]), _fmt_pct(r["pct_participacion"])])
    t = Table(data, repeatRows=1,
              colWidths=[26 * mm, 60 * mm, 15 * mm, 34 * mm, 20 * mm])
    t.setStyle(_base_table_style())
    return t


def _twr_pdf_table(twr: dict) -> Table:
    data = [["Mes", "R mes (Dietz)", "Benchmark", "Factor (1+R)", "TWR acumulado"]]
    for r in twr.get("rows", []):
        data.append([
            f"{r['month_name']} {r['year']}",
            _fmt_pct(r.get("portfolio_return")),
            _fmt_pct(r.get("benchmark_return")),
            "—" if r.get("factor") is None else f"{r['factor']:.5f}",
            _fmt_pct(r.get("cumulative_twr")),
        ])
    t = Table(data, repeatRows=1,
              colWidths=[36 * mm, 30 * mm, 28 * mm, 30 * mm, 32 * mm])
    t.setStyle(_base_table_style())
    return t
